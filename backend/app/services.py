from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO, StringIO
from typing import Iterable, Optional

import csv
import openpyxl
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from .database import (
    AuditLog,
    CaseSuite,
    CaseSuiteVersion,
    CaseTreeNode,
    CaseTreeNodeType,
    CaseLinkType,
    CaseSetSnapshot,
    CaseStatus,
    CaseStatusHistory,
    Issue,
    IssueCaseLink,
    IssueStatus,
    PlanStatus,
    RunCase,
    RunCaseHistory,
    RunStatus,
    SuiteCaseRef,
    SuiteVersionStatus,
    TestCase,
    TestPlan,
    TestRun,
    UserAccount,
    Version,
    VersionCaseStatus,
)
from .errors import AppError

ISSUE_GROUP_STATUSES: dict[str, set[IssueStatus]] = {
    "all": {
        IssueStatus.new,
        IssueStatus.assigned,
        IssueStatus.fixing,
        IssueStatus.to_verify,
        IssueStatus.verify_failed,
        IssueStatus.closed,
        IssueStatus.rejected,
        IssueStatus.duplicate,
        IssueStatus.invalid,
        IssueStatus.deferred,
    },
    "open": {IssueStatus.new, IssueStatus.assigned, IssueStatus.fixing, IssueStatus.verify_failed},
    "pending": {IssueStatus.to_verify},
    "closed": {IssueStatus.closed},
    "invalid": {IssueStatus.invalid, IssueStatus.rejected, IssueStatus.duplicate, IssueStatus.deferred},
}


def get_statuses_by_group(group: str) -> set[IssueStatus]:
    if group not in ISSUE_GROUP_STATUSES:
        raise AppError("VALIDATION_ERROR", f"未知状态组: {group}", status_code=422)
    return ISSUE_GROUP_STATUSES[group]


def find_version_by_key(db: Session, version_key: str) -> Version:
    version = db.query(Version).filter(Version.version_key == version_key).first()
    if not version:
        raise AppError("NOT_FOUND", f"版本不存在: {version_key}", status_code=404)
    return version


def find_user_by_user_id(db: Session, user_id: str) -> UserAccount:
    user = db.query(UserAccount).filter(UserAccount.user_id == user_id).first()
    if not user:
        raise AppError("NOT_FOUND", f"用户不存在: {user_id}", status_code=404)
    return user


def find_case_by_key(db: Session, case_key: str) -> TestCase:
    case = db.query(TestCase).filter(TestCase.case_key == case_key).first()
    if not case:
        raise AppError("NOT_FOUND", f"用例不存在: {case_key}", status_code=404)
    return case


def split_case_module_path(module_value: str | None) -> list[str]:
    if not module_value:
        return ["未分类"]
    parts = [item.strip() for item in str(module_value).replace("\\", "/").split("/") if item.strip()]
    return parts or ["未分类"]


def generate_case_tree_node_id(db: Session) -> str:
    count = db.query(func.count(CaseTreeNode.id)).scalar() or 0
    seq = int(count) + 1
    while True:
        node_id = f"NODE-{seq:04d}"
        exists = db.query(CaseTreeNode).filter(CaseTreeNode.node_id == node_id).first()
        if not exists:
            return node_id
        seq += 1


def _next_prefixed_key(db: Session, model, key_column, prefix: str, width: int = 4) -> str:
    count = db.query(func.count(model.id)).scalar() or 0
    seq = int(count) + 1
    while True:
        key = f"{prefix}-{seq:0{width}d}"
        exists = db.query(model).filter(key_column == key).first()
        if not exists:
            return key
        seq += 1


def generate_suite_key(db: Session) -> str:
    return _next_prefixed_key(db, CaseSuite, CaseSuite.suite_key, "SUITE")


def generate_suite_version_key(db: Session) -> str:
    return _next_prefixed_key(db, CaseSuiteVersion, CaseSuiteVersion.suite_version_key, "SV")


def generate_plan_key(db: Session) -> str:
    return _next_prefixed_key(db, TestPlan, TestPlan.plan_key, "PLAN")


def generate_run_key(db: Session) -> str:
    return _next_prefixed_key(db, TestRun, TestRun.run_key, "RUN")


def generate_run_case_key(db: Session) -> str:
    return _next_prefixed_key(db, RunCase, RunCase.run_case_key, "RC", width=6)


def find_suite_by_key(db: Session, suite_key: str) -> CaseSuite:
    suite = db.query(CaseSuite).filter(CaseSuite.suite_key == suite_key).first()
    if not suite:
        raise AppError("NOT_FOUND", f"用例集不存在: {suite_key}", status_code=404)
    return suite


def find_suite_version_by_key(db: Session, suite_version_key: str) -> CaseSuiteVersion:
    suite_version = db.query(CaseSuiteVersion).filter(CaseSuiteVersion.suite_version_key == suite_version_key).first()
    if not suite_version:
        raise AppError("NOT_FOUND", f"用例集版本不存在: {suite_version_key}", status_code=404)
    return suite_version


def find_plan_by_key(db: Session, plan_key: str) -> TestPlan:
    plan = db.query(TestPlan).filter(TestPlan.plan_key == plan_key).first()
    if not plan:
        raise AppError("NOT_FOUND", f"测试计划不存在: {plan_key}", status_code=404)
    return plan


def find_run_by_key(db: Session, run_key: str) -> TestRun:
    run = db.query(TestRun).filter(TestRun.run_key == run_key).first()
    if not run:
        raise AppError("NOT_FOUND", f"测试运行不存在: {run_key}", status_code=404)
    return run


def find_run_case_by_key(db: Session, run: TestRun, run_case_key: str) -> RunCase:
    run_case = db.query(RunCase).filter(RunCase.run_id == run.id, RunCase.run_case_key == run_case_key).first()
    if not run_case:
        raise AppError("NOT_FOUND", f"Run 用例不存在: {run_case_key}", status_code=404)
    return run_case


def find_case_tree_node_by_node_id(db: Session, node_id: str) -> CaseTreeNode:
    node = db.query(CaseTreeNode).filter(CaseTreeNode.node_id == node_id).first()
    if not node:
        raise AppError("NOT_FOUND", f"目录节点不存在: {node_id}", status_code=404)
    return node


def find_or_create_case_tree_node(
    db: Session,
    *,
    name: str,
    node_type: CaseTreeNodeType,
    parent: CaseTreeNode | None,
) -> CaseTreeNode:
    normalized_name = name.strip()
    if not normalized_name:
        raise AppError("VALIDATION_ERROR", "目录名不能为空", status_code=422)
    if parent and parent.node_type != CaseTreeNodeType.directory:
        raise AppError("VALIDATION_ERROR", "仅目录节点可新增子目录", status_code=422)

    full_path = normalized_name if not parent else f"{parent.full_path}/{normalized_name}"
    existing = db.query(CaseTreeNode).filter(CaseTreeNode.full_path == full_path).first()
    if existing:
        if existing.node_type != node_type:
            raise AppError("VALIDATION_ERROR", "同路径存在不同类型节点", status_code=422)
        return existing

    sibling = (
        db.query(CaseTreeNode)
        .filter(
            CaseTreeNode.parent_id == (parent.id if parent else None),
            CaseTreeNode.name == normalized_name,
        )
        .first()
    )
    if sibling:
        if sibling.node_type != node_type:
            raise AppError("VALIDATION_ERROR", "同级目录重名且类型冲突", status_code=422)
        return sibling

    node = CaseTreeNode(
        node_id=generate_case_tree_node_id(db),
        name=normalized_name,
        node_type=node_type,
        parent_id=parent.id if parent else None,
        full_path=full_path,
    )
    db.add(node)
    db.flush()
    return node


def ensure_case_tree_nodes_for_legacy_cases(db: Session) -> int:
    rows = db.query(TestCase).filter(TestCase.status == "active", TestCase.tree_node_id.is_(None)).all()
    if not rows:
        return 0

    updated = 0
    for test_case in rows:
        parent: CaseTreeNode | None = None
        for part in split_case_module_path(test_case.module):
            parent = find_or_create_case_tree_node(
                db,
                name=part,
                node_type=CaseTreeNodeType.directory,
                parent=parent,
            )

        file_node = find_or_create_case_tree_node(
            db,
            name="legacy_cases",
            node_type=CaseTreeNodeType.file,
            parent=parent,
        )
        test_case.tree_node_id = file_node.id
        updated += 1

    db.flush()
    return updated


def build_case_tree_rows(db: Session, version: Version) -> list[dict]:
    snapshot = get_or_create_default_snapshot(db, version)
    ensure_version_case_status_rows(db, version, snapshot, None)
    ensure_case_tree_nodes_for_legacy_cases(db)
    db.flush()

    status_rows = (
        db.query(VersionCaseStatus.case_id, VersionCaseStatus.status)
        .filter(
            VersionCaseStatus.version_id == version.id,
            VersionCaseStatus.snapshot_id == snapshot.id,
        )
        .all()
    )
    status_by_case_id = {row.case_id: row.status for row in status_rows}

    node_rows = db.query(CaseTreeNode).order_by(CaseTreeNode.full_path.asc()).all()
    case_rows = (
        db.query(TestCase)
        .filter(TestCase.status == "active")
        .order_by(TestCase.case_key.asc())
        .all()
    )

    node_payload: dict[int, dict] = {
        node.id: {
            "node_id": node.node_id,
            "name": node.name,
            "node_type": node.node_type,
            "parent_node_id": None,
            "full_path": node.full_path,
            "children": [],
            "cases": [],
        }
        for node in node_rows
    }

    roots: list[dict] = []
    for node in node_rows:
        payload = node_payload[node.id]
        if node.parent_id and node.parent_id in node_payload:
            payload["parent_node_id"] = node_payload[node.parent_id]["node_id"]
            node_payload[node.parent_id]["children"].append(payload)
        else:
            roots.append(payload)

    for test_case in case_rows:
        if not test_case.tree_node_id or test_case.tree_node_id not in node_payload:
            continue
        node_payload[test_case.tree_node_id]["cases"].append(
            {
                "case_key": test_case.case_key,
                "title": test_case.title,
                "module": test_case.module,
                "latest_status": status_by_case_id.get(test_case.id, CaseStatus.not_run),
            }
        )

    def sort_tree(node: dict) -> None:
        node["children"].sort(
            key=lambda item: (
                item["node_type"] != CaseTreeNodeType.directory,
                str(item["name"]).lower(),
            )
        )
        node["cases"].sort(key=lambda item: str(item["case_key"]).lower())
        for child in node["children"]:
            sort_tree(child)

    roots.sort(
        key=lambda item: (
            item["node_type"] != CaseTreeNodeType.directory,
            str(item["name"]).lower(),
        )
    )
    for root in roots:
        sort_tree(root)

    return roots


PLACEHOLDER_CASE_PATTERNS: dict[str, dict[str, set[str]]] = {
    "TC-001": {"modules": {"用户登录"}, "titles": {"用户登录成功"}},
    "TC-002": {"modules": {"用户登录"}, "titles": {"用户登录失败提示"}},
    "TC-003": {"modules": {"用户管理"}, "titles": {"新增用户成功", "用户创建成功"}},
    "TC-004": {"modules": {"用户管理", "数据导出"}, "titles": {"问题单导出 XLSX", "角色权限校验"}},
    "TC-005": {"modules": {"数据导出"}, "titles": {"数据导出 CSV"}},
    "TC-006": {"modules": {"数据导出"}, "titles": {"附件上传成功"}},
}


def cleanup_placeholder_cases(db: Session) -> dict[str, int]:
    candidates = (
        db.query(TestCase)
        .filter(TestCase.case_key.in_(list(PLACEHOLDER_CASE_PATTERNS.keys())))
        .all()
    )
    case_ids: list[int] = []
    for test_case in candidates:
        pattern = PLACEHOLDER_CASE_PATTERNS.get(test_case.case_key)
        if not pattern:
            continue
        if test_case.module in pattern["modules"] and test_case.title in pattern["titles"]:
            case_ids.append(test_case.id)

    if not case_ids:
        return {"deleted_cases": 0, "deleted_case_status": 0, "deleted_case_history": 0, "deleted_issue_links": 0}

    deleted_issue_links = (
        db.query(IssueCaseLink).filter(IssueCaseLink.case_id.in_(case_ids)).delete(synchronize_session=False)
    )
    deleted_case_status = (
        db.query(VersionCaseStatus).filter(VersionCaseStatus.case_id.in_(case_ids)).delete(synchronize_session=False)
    )
    deleted_case_history = (
        db.query(CaseStatusHistory).filter(CaseStatusHistory.case_id.in_(case_ids)).delete(synchronize_session=False)
    )
    deleted_cases = db.query(TestCase).filter(TestCase.id.in_(case_ids)).delete(synchronize_session=False)

    return {
        "deleted_cases": int(deleted_cases or 0),
        "deleted_case_status": int(deleted_case_status or 0),
        "deleted_case_history": int(deleted_case_history or 0),
        "deleted_issue_links": int(deleted_issue_links or 0),
    }


def get_or_create_default_snapshot(
    db: Session,
    version: Version,
    actor_id: Optional[int] = None,
) -> CaseSetSnapshot:
    existing = db.query(CaseSetSnapshot).filter(CaseSetSnapshot.version_id == version.id).first()
    if existing:
        return existing

    snapshot_count = db.query(func.count(CaseSetSnapshot.id)).scalar() or 0
    snapshot = CaseSetSnapshot(
        snapshot_id=f"SNAP-{snapshot_count + 1:03d}",
        version_id=version.id,
        case_set_id=None,
        snapshot_json={"scope": "all_active_cases", "reason": "auto-generated"},
        created_by=actor_id,
    )
    db.add(snapshot)
    db.flush()
    return snapshot


def ensure_version_case_status_rows(
    db: Session,
    version: Version,
    snapshot: CaseSetSnapshot,
    actor_id: Optional[int],
) -> None:
    cases = db.query(TestCase).filter(TestCase.status == "active").all()
    for test_case in cases:
        exists = (
            db.query(VersionCaseStatus)
            .filter(
                VersionCaseStatus.version_id == version.id,
                VersionCaseStatus.snapshot_id == snapshot.id,
                VersionCaseStatus.case_id == test_case.id,
            )
            .first()
        )
        if exists:
            continue

        db.add(
            VersionCaseStatus(
                version_id=version.id,
                snapshot_id=snapshot.id,
                case_id=test_case.id,
                status=CaseStatus.not_run,
                updated_by=actor_id,
            )
        )


def _suite_case_query(db: Session, suite_version_id: int):
    return (
        db.query(SuiteCaseRef, TestCase)
        .join(TestCase, TestCase.id == SuiteCaseRef.case_id)
        .filter(
            SuiteCaseRef.suite_version_id == suite_version_id,
            TestCase.status == "active",
        )
        .order_by(SuiteCaseRef.order_no.asc(), TestCase.case_key.asc())
    )


def _run_case_pass_rate(statuses: list[CaseStatus]) -> float | None:
    if not statuses:
        return None
    executed_statuses = {CaseStatus.passed, CaseStatus.failed, CaseStatus.blocked, CaseStatus.skipped}
    executed = len([item for item in statuses if item in executed_statuses])
    skipped = len([item for item in statuses if item == CaseStatus.skipped])
    passed = len([item for item in statuses if item == CaseStatus.passed])
    effective = executed - skipped
    if effective <= 0:
        return None
    return passed / effective


def calculate_run_metrics(db: Session, run: TestRun) -> dict:
    statuses = [row.current_status for row in db.query(RunCase).filter(RunCase.run_id == run.id).all()]
    total_cases = len(statuses)
    executed_cases = len(
        [
            item
            for item in statuses
            if item in {CaseStatus.passed, CaseStatus.failed, CaseStatus.blocked, CaseStatus.skipped}
        ]
    )
    return {
        "total_cases": total_cases,
        "executed_cases": executed_cases,
        "pass_rate": _run_case_pass_rate(statuses),
    }


def refresh_run_status(db: Session, run: TestRun) -> None:
    statuses = [row.current_status for row in db.query(RunCase).filter(RunCase.run_id == run.id).all()]
    if not statuses:
        run.status = RunStatus.running
        run.finished_at = None
        return

    if any(item == CaseStatus.not_run for item in statuses):
        run.status = RunStatus.running
        run.finished_at = None
        return

    if any(item in {CaseStatus.failed, CaseStatus.blocked} for item in statuses):
        run.status = RunStatus.failed
        run.finished_at = datetime.utcnow()
        return

    run.status = RunStatus.success
    run.finished_at = datetime.utcnow()


def create_run_from_suite_version(
    db: Session,
    *,
    plan: TestPlan,
    name: str,
    actor_id: int | None,
    build_no: str | None = None,
    environment: str | None = None,
    case_keys: set[str] | None = None,
) -> TestRun:
    run = TestRun(
        run_key=generate_run_key(db),
        plan_id=plan.id,
        name=name,
        build_no=build_no,
        environment=environment,
        status=RunStatus.running,
        created_by=actor_id,
    )
    db.add(run)
    db.flush()

    rows = _suite_case_query(db, plan.suite_version_id).all()
    for _, test_case in rows:
        if case_keys and test_case.case_key not in case_keys:
            continue
        run_case = RunCase(
            run_case_key=generate_run_case_key(db),
            run_id=run.id,
            case_id=test_case.id,
            case_key_snapshot=test_case.case_key,
            title_snapshot=test_case.title,
            module_snapshot=test_case.module,
            steps_snapshot=test_case.steps,
            expected_snapshot=test_case.expected,
            assignee_id=None,
            current_status=CaseStatus.not_run,
            last_updated_by=actor_id,
        )
        db.add(run_case)
        # Ensure subsequent key generation sees inserted rows and avoids duplicate run_case_key.
        db.flush()
    return run


def ensure_suite_case_refs_from_active_cases(
    db: Session,
    *,
    suite_version: CaseSuiteVersion,
) -> int:
    existing_case_ids = {
        row.case_id for row in db.query(SuiteCaseRef.case_id).filter(SuiteCaseRef.suite_version_id == suite_version.id).all()
    }
    active_cases = (
        db.query(TestCase)
        .filter(TestCase.status == "active")
        .order_by(TestCase.case_key.asc())
        .all()
    )
    order_no = db.query(func.count(SuiteCaseRef.id)).filter(SuiteCaseRef.suite_version_id == suite_version.id).scalar() or 0
    inserted = 0
    for test_case in active_cases:
        if test_case.id in existing_case_ids:
            continue
        order_no += 1
        db.add(
            SuiteCaseRef(
                suite_version_id=suite_version.id,
                case_id=test_case.id,
                order_no=order_no,
                module_path_snapshot=test_case.module,
            )
        )
        inserted += 1
    if inserted:
        db.flush()
    return inserted


def ensure_default_suite_version_for_version(
    db: Session,
    *,
    version: Version,
    actor_id: int | None,
) -> CaseSuiteVersion:
    suite = (
        db.query(CaseSuite)
        .filter(CaseSuite.version_id == version.id, CaseSuite.name == "默认用例集")
        .first()
    )
    if not suite:
        suite = CaseSuite(
            suite_key=generate_suite_key(db),
            version_id=version.id,
            name="默认用例集",
            description="系统自动维护的默认用例集",
            created_by=actor_id,
        )
        db.add(suite)
        db.flush()

    suite_version = (
        db.query(CaseSuiteVersion)
        .filter(CaseSuiteVersion.suite_id == suite.id, CaseSuiteVersion.ver_no == 1)
        .first()
    )
    if not suite_version:
        suite_version = CaseSuiteVersion(
            suite_version_key=generate_suite_version_key(db),
            suite_id=suite.id,
            ver_no=1,
            status=SuiteVersionStatus.published,
            note="默认发布版本",
            created_by=actor_id,
            published_by=actor_id,
            published_at=datetime.utcnow(),
        )
        db.add(suite_version)
        db.flush()

    ensure_suite_case_refs_from_active_cases(db, suite_version=suite_version)
    return suite_version


def _latest_legacy_status_map(db: Session, version_id: int) -> dict[int, tuple[CaseStatus, int | None, datetime]]:
    rows = (
        db.query(CaseStatusHistory)
        .filter(CaseStatusHistory.version_id == version_id)
        .order_by(CaseStatusHistory.executed_at.asc(), CaseStatusHistory.id.asc())
        .all()
    )
    status_map: dict[int, tuple[CaseStatus, int | None, datetime]] = {}
    for item in rows:
        status_map[item.case_id] = (item.status, item.executor_id, item.executed_at)
    return status_map


def ensure_legacy_plan_run_for_version(
    db: Session,
    *,
    version: Version,
    actor_id: int | None,
) -> TestRun:
    suite_version = ensure_default_suite_version_for_version(db, version=version, actor_id=actor_id)

    plan_name = f"历史迁移计划-{version.version_key}"
    plan = (
        db.query(TestPlan)
        .filter(TestPlan.version_id == version.id, TestPlan.name == plan_name)
        .first()
    )
    if not plan:
        plan = TestPlan(
            plan_key=generate_plan_key(db),
            version_id=version.id,
            suite_version_id=suite_version.id,
            name=plan_name,
            description="由旧版本执行状态迁移自动生成",
            status=PlanStatus.in_progress,
            created_by=actor_id,
        )
        db.add(plan)
        db.flush()

    run_name = f"Legacy-Run-{version.version_key}"
    run = db.query(TestRun).filter(TestRun.plan_id == plan.id, TestRun.name == run_name).first()
    if not run:
        run = TestRun(
            run_key=generate_run_key(db),
            plan_id=plan.id,
            name=run_name,
            status=RunStatus.running,
            created_by=actor_id,
        )
        db.add(run)
        db.flush()

    latest_status_map = _latest_legacy_status_map(db, version.id)
    default_snapshot = get_or_create_default_snapshot(db, version, actor_id)
    matrix_status_rows = (
        db.query(VersionCaseStatus)
        .filter(
            VersionCaseStatus.version_id == version.id,
            VersionCaseStatus.snapshot_id == default_snapshot.id,
        )
        .all()
    )
    matrix_status_map = {item.case_id: item for item in matrix_status_rows}

    existing_run_case_by_case_id = {
        row.case_id: row for row in db.query(RunCase).filter(RunCase.run_id == run.id).all()
    }

    for _, test_case in _suite_case_query(db, suite_version.id).all():
        run_case = existing_run_case_by_case_id.get(test_case.id)
        if not run_case:
            run_case = RunCase(
                run_case_key=generate_run_case_key(db),
                run_id=run.id,
                case_id=test_case.id,
                case_key_snapshot=test_case.case_key,
                title_snapshot=test_case.title,
                module_snapshot=test_case.module,
                steps_snapshot=test_case.steps,
                expected_snapshot=test_case.expected,
                assignee_id=None,
                current_status=CaseStatus.not_run,
                last_updated_by=actor_id,
            )
            db.add(run_case)
            db.flush()

        # 同步快照字段，避免后续库内容调整导致 Legacy Run 详情缺失
        run_case.case_key_snapshot = test_case.case_key
        run_case.title_snapshot = test_case.title
        run_case.module_snapshot = test_case.module
        run_case.steps_snapshot = test_case.steps
        run_case.expected_snapshot = test_case.expected

        if db.query(func.count(RunCaseHistory.id)).filter(RunCaseHistory.run_case_id == run_case.id).scalar() == 0:
            legacy_rows = (
                db.query(CaseStatusHistory)
                .filter(
                    CaseStatusHistory.version_id == version.id,
                    CaseStatusHistory.case_id == test_case.id,
                )
                .order_by(CaseStatusHistory.executed_at.asc(), CaseStatusHistory.id.asc())
                .all()
            )
            for legacy in legacy_rows:
                db.add(
                    RunCaseHistory(
                        run_case_id=run_case.id,
                        status=legacy.status,
                        remark=legacy.note,
                        operator_id=legacy.executor_id,
                        operated_at=legacy.executed_at,
                        attachments_json=legacy.attachments_json,
                    )
                )

        latest_legacy = latest_status_map.get(test_case.id)
        if latest_legacy:
            run_case.current_status = latest_legacy[0]
            run_case.last_updated_by = latest_legacy[1]
            run_case.last_updated_at = latest_legacy[2]
        else:
            matrix_row = matrix_status_map.get(test_case.id)
            run_case.current_status = matrix_row.status if matrix_row else CaseStatus.not_run
            run_case.last_updated_by = matrix_row.updated_by if matrix_row else actor_id
            run_case.last_updated_at = matrix_row.updated_at if matrix_row else datetime.utcnow()

    refresh_run_status(db, run)
    db.flush()
    return run


def ensure_legacy_plan_runs(db: Session, actor_id: int | None) -> None:
    versions = db.query(Version).all()
    for version in versions:
        ensure_legacy_plan_run_for_version(db, version=version, actor_id=actor_id)


def write_audit(
    db: Session,
    actor_id: Optional[int],
    action: str,
    object_type: str,
    object_id: str,
    diff: dict,
) -> None:
    db.add(
        AuditLog(
            actor_id=actor_id,
            action=action,
            object_type=object_type,
            object_id=object_id,
            diff_json=diff,
        )
    )


def calculate_dashboard(db: Session, version: Version) -> dict:
    latest_completed_run = (
        db.query(TestRun)
        .join(TestPlan, TestPlan.id == TestRun.plan_id)
        .filter(
            TestPlan.version_id == version.id,
            TestRun.status.in_([RunStatus.success, RunStatus.failed]),
        )
        .order_by(TestRun.finished_at.desc(), TestRun.started_at.desc())
        .first()
    )

    if latest_completed_run:
        statuses = [row.current_status for row in db.query(RunCase).filter(RunCase.run_id == latest_completed_run.id).all()]
        total_cases = len(statuses)
        executed_statuses = {CaseStatus.passed, CaseStatus.failed, CaseStatus.blocked, CaseStatus.skipped}
        executed_cases = len([item for item in statuses if item in executed_statuses])
        passed_cases = len([item for item in statuses if item == CaseStatus.passed])
        skipped_cases = len([item for item in statuses if item == CaseStatus.skipped])
        execution_rate = 0.0 if total_cases == 0 else executed_cases / total_cases
        effective_denominator = executed_cases - skipped_cases
        pass_rate = None if effective_denominator <= 0 else passed_cases / effective_denominator
    else:
        snapshot = get_or_create_default_snapshot(db, version)
        rows = (
            db.query(VersionCaseStatus, TestCase)
            .join(TestCase, TestCase.id == VersionCaseStatus.case_id)
            .filter(
                VersionCaseStatus.version_id == version.id,
                VersionCaseStatus.snapshot_id == snapshot.id,
                TestCase.status == "active",
            )
            .all()
        )

        total_cases = len(rows)
        executed_statuses = {CaseStatus.passed, CaseStatus.failed, CaseStatus.blocked, CaseStatus.skipped}
        executed_cases = len([1 for row, _ in rows if row.status in executed_statuses])
        passed_cases = len([1 for row, _ in rows if row.status == CaseStatus.passed])
        skipped_cases = len([1 for row, _ in rows if row.status == CaseStatus.skipped])
        execution_rate = 0.0 if total_cases == 0 else executed_cases / total_cases

        effective_denominator = executed_cases - skipped_cases
        pass_rate = None if effective_denominator <= 0 else passed_cases / effective_denominator

    issues = db.query(Issue).filter(Issue.found_version_id == version.id).all()
    total_issues = len(issues)
    resolved_statuses = ISSUE_GROUP_STATUSES["closed"] | ISSUE_GROUP_STATUSES["invalid"]
    resolved_issues = len([1 for issue in issues if issue.status in resolved_statuses])
    pending_issue_count = len([1 for issue in issues if issue.status in ISSUE_GROUP_STATUSES["pending"]])
    resolution_rate = 0.0 if total_issues == 0 else resolved_issues / total_issues

    return {
        "total_cases": total_cases,
        "executed_cases": executed_cases,
        "passed_cases": passed_cases,
        "skipped_cases": skipped_cases,
        "execution_rate": execution_rate,
        "pass_rate": pass_rate,
        "total_issues": total_issues,
        "resolved_issues": resolved_issues,
        "resolution_rate": resolution_rate,
        "pending_issue_count": pending_issue_count,
    }


def serialize_issue_links(links: Iterable[IssueCaseLink], db: Session) -> list[dict]:
    case_map = {
        row.id: row.case_key
        for row in db.query(TestCase).filter(TestCase.id.in_([link.case_id for link in links])).all()
    }
    return [
        {
            "case_key": case_map.get(link.case_id, ""),
            "link_type": link.link_type,
            "note": link.note,
        }
        for link in links
    ]


def issue_to_dict(issue: Issue, db: Session) -> dict:
    version_ids = [issue.found_version_id, issue.fix_version_id, issue.verify_version_id]
    version_map = {
        row.id: row.version_key
        for row in db.query(Version).filter(Version.id.in_([version_id for version_id in version_ids if version_id])).all()
    }

    user_ids = [issue.assignee_id, issue.reporter_id]
    user_map = {
        row.id: row.user_id
        for row in db.query(UserAccount).filter(UserAccount.id.in_([user_id for user_id in user_ids if user_id])).all()
    }

    return {
        "issue_key": issue.issue_key,
        "title": issue.title,
        "description": issue.description,
        "status": issue.status,
        "priority": issue.priority,
        "severity": issue.severity,
        "assignee": user_map.get(issue.assignee_id),
        "reporter": user_map.get(issue.reporter_id, ""),
        "found_version_key": version_map.get(issue.found_version_id, ""),
        "fix_version_key": version_map.get(issue.fix_version_id) if issue.fix_version_id else None,
        "verify_version_key": version_map.get(issue.verify_version_id) if issue.verify_version_id else None,
        "created_at": issue.created_at,
        "updated_at": issue.updated_at,
        "links": serialize_issue_links(issue.links, db),
    }


def as_csv_response(filename: str, headers: list[str], rows: list[dict]) -> StreamingResponse:
    text_stream = StringIO()
    writer = csv.DictWriter(text_stream, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        normalized = {key: row.get(key, "") for key in headers}
        writer.writerow(normalized)
    stream = BytesIO(text_stream.getvalue().encode("utf-8"))
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def as_xlsx_response(filename: str, headers: list[str], rows: list[dict]) -> StreamingResponse:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(key, "") for key in headers])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def parse_import_rows(filename: str, file_bytes: bytes) -> list[dict[str, str]]:
    lower = filename.lower()

    if lower.endswith(".csv"):
        text = file_bytes.decode("utf-8")
        reader = csv.DictReader(text.splitlines())
        return [
            {str(key).strip(): str(value).strip() for key, value in row.items() if key is not None}
            for row in reader
        ]

    if lower.endswith(".xlsx"):
        workbook = openpyxl.load_workbook(BytesIO(file_bytes))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() for value in rows[0]]
        result: list[dict[str, str]] = []
        for data_row in rows[1:]:
            mapped = {
                headers[idx]: str(value).strip() if value is not None else ""
                for idx, value in enumerate(data_row)
                if idx < len(headers)
            }
            if any(mapped.values()):
                result.append(mapped)
        return result

    raise AppError("VALIDATION_ERROR", "仅支持 CSV/XLSX 导入", status_code=422)


def validate_required_columns(rows: list[dict[str, str]], required_columns: list[str]) -> list[str]:
    errors: list[str] = []
    for idx, row in enumerate(rows):
        for key in required_columns:
            if not row.get(key):
                errors.append(f"第 {idx + 2} 行缺少 {key}")
    return errors


def build_case_module_rows(db: Session, version: Version) -> list[dict]:
    snapshot = get_or_create_default_snapshot(db, version)
    statuses = (
        db.query(VersionCaseStatus, TestCase, UserAccount)
        .join(TestCase, TestCase.id == VersionCaseStatus.case_id)
        .outerjoin(UserAccount, UserAccount.id == VersionCaseStatus.updated_by)
        .filter(
            VersionCaseStatus.version_id == version.id,
            VersionCaseStatus.snapshot_id == snapshot.id,
            TestCase.status == "active",
        )
        .all()
    )

    grouped: dict[str, list[dict]] = defaultdict(list)
    for status, test_case, updater in statuses:
        grouped[test_case.module].append(
            {
                "case_key": test_case.case_key,
                "title": test_case.title,
                "module": test_case.module,
                "latest_status": status.status,
                "latest_updated_at": status.updated_at,
                "latest_updated_by": updater.user_id if updater else None,
            }
        )

    return [
        {
            "module": module,
            "cases": sorted(items, key=lambda row: row["case_key"]),
        }
        for module, items in sorted(grouped.items(), key=lambda item: item[0])
    ]


def build_case_history_rows(db: Session, version_id: int, case_id: int, limit: int = 10) -> list[dict]:
    rows = (
        db.query(CaseStatusHistory, UserAccount)
        .outerjoin(UserAccount, UserAccount.id == CaseStatusHistory.executor_id)
        .filter(CaseStatusHistory.version_id == version_id, CaseStatusHistory.case_id == case_id)
        .order_by(CaseStatusHistory.executed_at.desc())
        .limit(limit)
        .all()
    )

    return [
        {
            "id": history.id,
            "executed_at": history.executed_at,
            "executor": user.user_id if user else "system",
            "status": history.status,
            "note": history.note,
            "attachments": history.attachments_json,
        }
        for history, user in rows
    ]


def make_snapshot_payload(version_key: str, case_keys: list[str], reason: str) -> dict:
    return {
        "version_key": version_key,
        "case_keys": case_keys,
        "reason": reason,
        "generated_at": datetime.utcnow().isoformat(),
    }
